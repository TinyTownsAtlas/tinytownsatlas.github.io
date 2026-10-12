"""Build the Atlas hospital-services context extract from the archived PHIDU workbooks.

PHIDU (Torrens University Australia) supplied hospital estimates by UCL that were produced from
Population Health Area (PHA) data: each UCL's share of its PHA's population was applied to the
PHA's observed and expected admissions, and the standardised ratio recalculated. Many UCLs in one
PHA therefore carry identical rates and ratios. They are broader-area context, not observations
for the town. See `data_sources_notes-final.docx` ("Geographical estimation, presentation and
caveats").

Indicators extracted (indirectly age-standardised ratio, Australia = 100):
  - public_admissions: "Hospital admissions - Public hospitals" / "Total admissions", 2022/23
    (excluding same-day admissions for renal dialysis), workbook columns 2-4.
  - pph: "Potentially preventable hospitalisations (PPH) - Public hospitals" / "PPH admissions for
    all conditions", 2020/21, workbook columns 21-23.
Allocated admission numbers are deliberately not extracted.

Source symbols (workbook "Key" sheet): "#" not shown (population under 100, or fewer than 6
admissions); ".." not applicable; "n.a." not available. A row whose every numeric cell is 0
(all admissions, PPH and ED measures) is the pattern of the no-resident "Migratory - Offshore -
Shipping" records and is read as "no data assigned", not as observed zero admissions: it is kept as
"not_interpretable". In the Atlas universe this is Uriarra Village (L), UCL822002. Any other zero
ratio stops the build for review rather than being withheld.

Inputs (read-only):
  05_source_archive/PHIDU_John_Glover/prior_validated_sources/Hospital_data_workbook.xlsx
  05_source_archive/PHIDU_John_Glover/prior_validated_sources/data_sources_notes-final.docx
  05_source_archive/PHIDU_John_Glover/2026-09-23_final_delivery/_data_workbook-witth_details.xlsx
  06_derived_data/census/atlas_ucl_dataset.csv          (Atlas UCL universe)

Outputs:
  06_derived_data/phidu_context/atlas_hospital_context.csv
  06_derived_data/phidu_context/atlas_hospital_context_provenance.json

Run after build_analytic.py and before build_data.py.
"""

from __future__ import annotations

import hashlib
import json
import math
from collections import defaultdict
from datetime import datetime, timezone
from pathlib import Path

import re

import openpyxl
import pandas as pd

ROOT = Path(__file__).resolve().parents[2]
PHIDU_DIR = ROOT / "05_source_archive" / "PHIDU_John_Glover"
HOSPITAL_WORKBOOK = PHIDU_DIR / "prior_validated_sources" / "Hospital_data_workbook.xlsx"
NOTES_DOCX = PHIDU_DIR / "prior_validated_sources" / "data_sources_notes-final.docx"
MASTER_WORKBOOK = PHIDU_DIR / "2026-09-23_final_delivery" / "_data_workbook-witth_details.xlsx"
ATLAS_CSV = ROOT / "06_derived_data" / "census" / "atlas_ucl_dataset.csv"
OUT_DIR = ROOT / "06_derived_data" / "phidu_context"
OUT_CSV = OUT_DIR / "atlas_hospital_context.csv"
OUT_PROVENANCE = OUT_DIR / "atlas_hospital_context_provenance.json"

# sha256 recorded in prior_validated_sources/diabetology_source_manifest.csv when the workbook was
# first validated; the build stops if the archived file has changed.
HOSPITAL_WORKBOOK_SHA256 = "6dfcf979d400641b3c88ff2a33d51fb7a5d05326f96df7b19c100b51651a2473"

# Indicator -> (first column of its Number/Rate/SR triple, block label (row 0), measure label (row 1)).
INDICATORS = {
    "public_admissions": (2, "Hospital admissions - Public hospitals", "Total admissions"),
    "pph": (21, "Potentially preventable hospitalisations (PPH) - Public hospitals", "PPH admissions for all conditions"),
}
TRIPLE = ("Number", "Rate", "SR")

# Master workbook PHA flag columns (row 1 block label, row 2 headers).
PHA_SECTION = (462, "UCL Flag - Hospital and Mortality data")
PHA_CODE_COL = (462, "PHA_CODE_2021")
PHA_NAME_COL = (463, "PHA_NAME")

UCL_CODE = re.compile(r"^UCL\d{6}$")

SYMBOL_STATUS = {"#": "suppressed", "..": "suppressed", "n.a.": "not_available"}


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def label(value: object) -> str:
    return str(value or "").strip()


def read_hospital() -> tuple[dict[str, tuple], dict[str, tuple], set[str]]:
    workbook = openpyxl.load_workbook(HOSPITAL_WORKBOOK, read_only=True, data_only=True)
    try:
        rows = list(workbook["Data"].iter_rows(values_only=True))
        aggregation_rows = list(workbook["UCL_aggregations"].iter_rows(values_only=True))
    finally:
        workbook.close()

    blocks, measures, headers = rows[0], rows[1], rows[2]
    if (label(headers[0]), label(headers[1])) != ("UCL_CODE_2021", "UCL_NAME"):
        raise SystemExit("Hospital workbook: unexpected identifier headers")
    for key, (col, block, measure) in INDICATORS.items():
        if label(blocks[col]) != block or label(measures[col]) != measure:
            raise SystemExit(f"Hospital workbook column {col}: expected {block!r} / {measure!r}, found "
                             f"{label(blocks[col])!r} / {label(measures[col])!r}")
        if tuple(label(h) for h in headers[col:col + 3]) != TRIPLE:
            raise SystemExit(f"Hospital workbook column {col}: expected Number/Rate/SR headers for {key}")

    data: dict[str, tuple] = {}
    totals: dict[str, tuple] = {}
    for row in rows[3:]:
        code = row[0]
        if isinstance(code, str) and UCL_CODE.match(code):
            if code in data:
                raise SystemExit(f"Duplicate UCL code in hospital workbook: {code}")
            data[code] = row
        elif code == "Total":
            totals["Total"] = row
    split_codes = {r[0] for r in aggregation_rows if isinstance(r[0], str) and UCL_CODE.match(r[0])}
    return data, totals, split_codes


def read_master_pha() -> dict[str, tuple[str, str]]:
    workbook = openpyxl.load_workbook(MASTER_WORKBOOK, read_only=True, data_only=True)
    try:
        rows = workbook["data"].iter_rows(values_only=True)
        sections, header = next(rows), next(rows)
        if not label(sections[PHA_SECTION[0]]).startswith(PHA_SECTION[1]):
            raise SystemExit(f"Master workbook column {PHA_SECTION[0]}: expected block {PHA_SECTION[1]!r}")
        for col, expected in (PHA_CODE_COL, PHA_NAME_COL):
            if label(header[col]) != expected:
                raise SystemExit(f"Master workbook column {col}: expected {expected!r}, found {label(header[col])!r}")
        out = {}
        for row in rows:
            code = row[0]
            if isinstance(code, str) and UCL_CODE.match(code):
                pha_code, pha_name = label(row[PHA_CODE_COL[0]]), label(row[PHA_NAME_COL[0]])
                out[code] = ("", "") if pha_code in ("", "..") else (pha_code, pha_name)
    finally:
        workbook.close()
    return out


def all_zero_row(row: tuple) -> bool:
    numeric = [v for v in row[2:] if isinstance(v, (int, float)) and not isinstance(v, bool)]
    return bool(numeric) and all(v == 0 for v in numeric)


def parse_value(raw: object, code: str, column: int, row_all_zero: bool) -> tuple[float | None, str]:
    if isinstance(raw, bool):
        raise SystemExit(f"{code}: hospital workbook column {column} has a boolean value")
    if isinstance(raw, (int, float)):
        if not math.isfinite(raw) or raw < 0:
            raise SystemExit(f"{code}: hospital workbook column {column} has invalid value {raw!r}")
        if raw == 0:
            if row_all_zero:
                return None, "not_interpretable"
            raise SystemExit(f"{code}: zero value in hospital workbook column {column} outside an all-zero row; review before publishing")
        return float(raw), "available"
    symbol = label(raw)
    if symbol not in SYMBOL_STATUS:
        raise SystemExit(f"{code}: hospital workbook column {column} has unknown symbol {raw!r}")
    return None, SYMBOL_STATUS[symbol]


def main() -> None:
    actual_hash = sha256(HOSPITAL_WORKBOOK)
    if actual_hash != HOSPITAL_WORKBOOK_SHA256:
        raise SystemExit(f"Hospital workbook hash {actual_hash} differs from the validated copy")

    data, totals, split_codes = read_hospital()
    pha = read_master_pha()
    atlas = pd.read_csv(ATLAS_CSV, dtype={"ucl_code": str})

    if set(data) != set(pha):
        raise SystemExit("Hospital workbook and master workbook UCL code sets differ")
    flagged = {code for code, (pha_code, _) in pha.items() if pha_code}
    if flagged != split_codes:
        raise SystemExit("Master PHA flags do not match the hospital workbook's UCL_aggregations sheet")
    for key, (col, _, _) in INDICATORS.items():
        national_sr = totals["Total"][col + 2]
        if not abs(float(national_sr) - 100) < 1e-6:
            raise SystemExit(f"{key}: national total SR is {national_sr}, expected 100")

    # Every UCL drawn from the same split PHA must carry the same ratio (PHIDU's allocation method).
    by_pha: dict[str, list[str]] = defaultdict(list)
    for code, (pha_code, _) in pha.items():
        if pha_code:
            by_pha[pha_code].append(code)
    for key, (col, _, _) in INDICATORS.items():
        for pha_code, members in by_pha.items():
            values = {round(float(data[c][col + 2]), 6) for c in members if isinstance(data[c][col + 2], (int, float))}
            if len(values) > 1:
                raise SystemExit(f"{key}: UCLs from PHA {pha_code} carry different ratios {sorted(values)}")

    # Which UCLs share one area's values. The PHA flag lists only UCLs with under 50% of their
    # PHA's population; the PHA's main UCL is unflagged but carries the same ratios (e.g. Broken
    # Hill with Wilcannia, Ivanhoe and Menindee). So sharing is counted from the values: UCLs with
    # an identical public-admissions ratio (never suppressed for real areas; matches are exact to
    # ~1e-12). Where two such UCLs both have a published PPH ratio, it must match too, and every
    # flagged PHA group must fall inside one sharing set.
    adm_col = INDICATORS["public_admissions"][0] + 2
    pph_col = INDICATORS["pph"][0] + 2
    sharing: dict[float, list[str]] = defaultdict(list)
    for code, row in data.items():
        if isinstance(row[adm_col], (int, float)) and row[adm_col] > 0:
            sharing[round(float(row[adm_col]), 6)].append(code)
    for members in sharing.values():
        pph_values = {round(float(data[c][pph_col]), 6) for c in members if isinstance(data[c][pph_col], (int, float)) and data[c][pph_col] > 0}
        if len(pph_values) > 1:
            raise SystemExit(f"UCLs {members} share an admissions ratio but not a PPH ratio")
    for pha_code, members in by_pha.items():
        keys = {round(float(data[c][adm_col]), 6) for c in members if isinstance(data[c][adm_col], (int, float)) and data[c][adm_col] > 0}
        if len(keys) > 1:
            raise SystemExit(f"PHA {pha_code} members fall in different sharing sets")

    def shared_count(code: str) -> int:
        value = data[code][adm_col]
        if isinstance(value, (int, float)) and value > 0:
            return len(sharing[round(float(value), 6)])
        return 1

    records = []
    for code in atlas["ucl_code"]:
        row = data[code]
        if str(row[1]) != atlas.loc[atlas["ucl_code"] == code, "ucl_name"].item():
            raise SystemExit(f"{code}: UCL name differs between hospital workbook and Atlas dataset")
        pha_code, pha_name = pha[code]
        rec: dict[str, object] = {
            "ucl_code": code,
            "pha_code": pha_code,
            "pha_name": pha_name,
            "pha_split": bool(pha_code),
            # UCLs in the whole PHIDU workbook (any size) with this UCL's values, including itself.
            "shared_ucl_count": shared_count(code),
        }
        for key, (col, _, _) in INDICATORS.items():
            sr, status = parse_value(row[col + 2], code, col + 2, all_zero_row(row))
            rate, rate_status = parse_value(row[col + 1], code, col + 1, all_zero_row(row))
            if status != rate_status:
                raise SystemExit(f"{code} {key}: rate status {rate_status} differs from ratio status {status}")
            rec[f"{key}_status"] = status
            rec[f"{key}_sr"] = sr
            rec[f"{key}_rate_per_100k"] = rate
        records.append(rec)

    out = pd.DataFrame(records)
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    out.to_csv(OUT_CSV, index=False)

    status_counts = {
        key: out[f"{key}_status"].value_counts().sort_index().to_dict() for key in INDICATORS
    }
    provenance = {
        "generatedAt": datetime.now(timezone.utc).isoformat(timespec="seconds"),
        "script": "04_site/scripts/build_phidu_context.py",
        "sources": [
            {"path": str(p.relative_to(ROOT)).replace("\\", "/"), "sha256": sha256(p)}
            for p in (HOSPITAL_WORKBOOK, NOTES_DOCX, MASTER_WORKBOOK, ATLAS_CSV)
        ],
        "licence": "CC BY-NC-SA 3.0 AU (PHIDU content; stated in data_sources_notes-final.docx)",
        "indicators": {
            "public_admissions": "Total public hospital admissions (excluding same-day renal dialysis), 2022/23; indirectly age-standardised ratio, Australia = 100",
            "pph": "Potentially preventable hospitalisations, all conditions, public hospitals, 2020/21; indirectly age-standardised ratio, Australia = 100",
        },
        "rules": {
            "geography": "PHA-derived estimates allocated to UCLs by PHIDU; pha_code is blank where the UCL is not in PHIDU's UCL_aggregations list (UCLs split from a PHA with under 50% of its population)",
            "symbols": "'#' and '..' -> suppressed; 'n.a.' -> not_available; a row with every numeric cell 0 (same pattern as the no-resident Migratory records) -> not_interpretable; any other zero ratio stops the build",
            "omitted": "Allocated admission numbers are not extracted",
        },
        "counts": {"atlasUcls": len(out), "phaSplitUcls": int(out["pha_split"].sum()), "status": status_counts},
    }
    OUT_PROVENANCE.write_text(json.dumps(provenance, indent=2), encoding="utf-8")

    print(f"Wrote {len(out)} rows to {OUT_CSV}")
    print(f"UCLs with a split PHA: {int(out['pha_split'].sum())}")
    for key, counts in status_counts.items():
        print(f"{key}: {counts}")


if __name__ == "__main__":
    main()
